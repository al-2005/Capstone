from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt

OUTPUT = Path(__file__).with_name("presentation.pptx")
SLIDE_W, SLIDE_H = 13.333, 7.5
NAVY = RGBColor(23, 34, 63)
BLUE = RGBColor(73, 108, 225)
PALE = RGBColor(239, 243, 252)
INK = RGBColor(42, 49, 67)
MUTED = RGBColor(112, 120, 139)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(71, 166, 119)
ORANGE = RGBColor(224, 151, 66)

presentation = Presentation()
presentation.slide_width = Inches(SLIDE_W)
presentation.slide_height = Inches(SLIDE_H)
blank = presentation.slide_layouts[6]


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False, font="Aptos"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(8)
    return shape


def base_slide(title, subtitle=None):
    slide = presentation.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(1.22)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    textbox(slide, 0.62, 0.2, 12, 0.52, title, 25, WHITE, True, "Aptos Display")
    if subtitle:
        textbox(slide, 0.65, 0.78, 12, 0.27, subtitle, 10, RGBColor(202, 211, 231))
    textbox(slide, 0.65, 7.12, 10, 0.2, "CAMPUS HELPDESK  ·  DEVSECOPS CAPSTONE", 8, MUTED, True)
    textbox(slide, 12.15, 7.1, 0.55, 0.2, str(len(presentation.slides)), 9, MUTED)
    return slide


def bullet_list(slide, items, x=0.9, y=1.65, w=11.5, size=19, gap=0.73):
    for index, item in enumerate(items):
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y + index * gap + 0.1), Inches(0.12), Inches(0.12)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = BLUE
        marker.line.fill.background()
        textbox(slide, x + 0.3, y + index * gap, w - 0.3, gap, item, size, INK)


def card(slide, x, y, w, h, title, detail, accent=BLUE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.color.rgb = RGBColor(222, 228, 241)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.04), Inches(0.08), Inches(h - 0.08)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    textbox(slide, x + 0.22, y + 0.18, w - 0.38, 0.38, title, 15, INK, True)
    textbox(slide, x + 0.22, y + 0.68, w - 0.38, h - 0.78, detail, 11, MUTED)


def connector(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    line.line.color.rgb = BLUE
    line.line.width = Pt(2)


slide = presentation.slides.add_slide(blank)
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = NAVY
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.05), Inches(0.12), Inches(4.6))
accent.fill.solid()
accent.fill.fore_color.rgb = BLUE
accent.line.fill.background()
textbox(slide, 1.15, 1.25, 10.8, 0.4, "DEVSECOPS CAPSTONE  ·  SESSION 21", 13, RGBColor(180, 197, 237), True)
textbox(slide, 1.12, 2.02, 10.8, 1.12, "Campus Helpdesk", 43, WHITE, True, "Aptos Display")
textbox(slide, 1.17, 3.34, 10.4, 0.8, "From a support request to a monitored cloud service", 21, RGBColor(216, 224, 242))
textbox(slide, 1.18, 5.72, 9.8, 0.8, "React  ·  FastAPI  ·  PostgreSQL  ·  Docker  ·  AWS EKS  ·  Argo CD  ·  Prometheus", 13, RGBColor(191, 204, 231))

slide = base_slide("01  |  The problem and the product", "A realistic campus service desk—not a generic task CRUD demo")
card(slide, 0.75, 1.7, 3.8, 2.0, "Students", "Report classroom, IT, facilities, and account issues with a clear priority.", BLUE)
card(slide, 4.78, 1.7, 3.8, 2.0, "Service staff", "Triage requests, assign an owner, update status, and close resolved tickets.", GREEN)
card(slide, 8.81, 1.7, 3.8, 2.0, "Service leads", "See queue volume, urgent work, aging requests, and service health.", ORANGE)
textbox(slide, 0.95, 4.35, 11.2, 1.35, "Goal: connect the application workflow to the engineering workflow—tests, secure image delivery, infrastructure as code, GitOps deployment, and observable operations.", 21, INK, True)

slide = base_slide("02  |  Application workflow", "The support lifecycle is visible in both the UI and REST API")
card(slide, 0.78, 1.72, 2.72, 1.55, "1 · Submit", "Category, description, requester, and priority.", BLUE)
card(slide, 3.65, 1.72, 2.72, 1.55, "2 · Triage", "Search the queue; classify and assign work.", ORANGE)
card(slide, 6.52, 1.72, 2.72, 1.55, "3 · Resolve", "Move through open, in progress, waiting, resolved.", GREEN)
card(slide, 9.39, 1.72, 2.72, 1.55, "4 · Learn", "Monitor ticket counts, SLA age, and API health.", BLUE)
connector(slide, 3.5, 2.5, 3.63, 2.5)
connector(slide, 6.37, 2.5, 6.5, 2.5)
connector(slide, 9.24, 2.5, 9.37, 2.5)
bullet_list(slide, [
    "REST API supports ticket list, filters/search, create, read, partial update, delete, and stats.",
    "PostgreSQL stores durable tickets; Alembic manages schema changes.",
    "FastAPI validates request shape and publishes OpenAPI documentation."
], y=4.1, size=16, gap=0.62)

slide = base_slide("03  |  Architecture", "One product, one deployable path, explicit service boundaries")
card(slide, 0.55, 1.55, 2.18, 1.2, "Browser", "React dashboard", BLUE)
card(slide, 3.15, 1.55, 2.25, 1.2, "Ingress / NGINX", "Route / and /api", ORANGE)
card(slide, 5.85, 1.55, 2.25, 1.2, "FastAPI", "REST + health + metrics", GREEN)
card(slide, 8.55, 1.55, 2.0, 1.2, "PostgreSQL", "Persistent ticket store", BLUE)
card(slide, 10.9, 1.55, 1.85, 1.2, "Monitoring", "Prometheus + Grafana", ORANGE)
connector(slide, 2.75, 2.15, 3.12, 2.15)
connector(slide, 5.4, 2.15, 5.82, 2.15)
connector(slide, 8.1, 2.15, 8.52, 2.15)
connector(slide, 8.0, 2.78, 10.98, 3.15)
card(slide, 0.85, 4.15, 3.55, 1.45, "GitHub Actions", "pytest → Vite build → container build → Trivy → GHCR SHA tags", BLUE)
card(slide, 4.9, 4.15, 3.55, 1.45, "Terraform", "AWS VPC, subnets, NAT, EKS, and managed workers", ORANGE)
card(slide, 8.95, 4.15, 3.55, 1.45, "Argo CD + Helm", "Git is desired state; Argo CD reconciles the Kubernetes release", GREEN)
connector(slide, 4.4, 4.88, 4.88, 4.88)
connector(slide, 8.45, 4.88, 8.93, 4.88)

slide = base_slide("04  |  Technology and repository", "The folder structure maps to the system lifecycle")
card(slide, 0.8, 1.6, 5.65, 1.45, "Application", "frontend/ React + Vite · backend/ FastAPI + SQLAlchemy · PostgreSQL + Alembic", BLUE)
card(slide, 6.85, 1.6, 5.65, 1.45, "Quality and delivery", ".github/workflows/ CI · pytest · npm build · Dockerfiles · Trivy · GHCR", GREEN)
card(slide, 0.8, 3.38, 5.65, 1.45, "Infrastructure and runtime", "terraform/ AWS · helm/campus-helpdesk/ chart · k8s/ namespace · argocd/ Application", ORANGE)
card(slide, 6.85, 3.38, 5.65, 1.45, "Operations and evidence", "monitoring/ Prometheus + Grafana · troubleshooting/ labs · docs/screenshots/ · demo/", BLUE)
textbox(slide, 1.0, 5.55, 11, 0.72, "The README is organized by course module (M1–M10) so reviewers can connect each claim to commands, implementation files, and evidence.", 17, INK, True)

slide = base_slide("05  |  CI/CD and container security", "Tests and security checks are gates—not decoration")
card(slide, 0.78, 1.65, 2.65, 1.55, "Pull request", "Run API tests and frontend production build.", BLUE)
card(slide, 3.6, 1.65, 2.65, 1.55, "Build images", "Backend and multi-stage, non-root frontend.", ORANGE)
card(slide, 6.42, 1.65, 2.65, 1.55, "Trivy scan", "Block HIGH / CRITICAL vulnerabilities.", GREEN)
card(slide, 9.24, 1.65, 2.65, 1.55, "Publish", "GHCR packages tagged by commit SHA.", BLUE)
connector(slide, 3.43, 2.43, 3.58, 2.43)
connector(slide, 6.25, 2.43, 6.4, 2.43)
connector(slide, 9.07, 2.43, 9.22, 2.43)
bullet_list(slide, [
    "Fail fast on test/build errors: an untested revision does not become a release image.",
    "Use immutable tags to trace each runtime image back to a source commit.",
    "Argo CD watches the deployment configuration and reports sync/health separately."
], y=4.05, size=16, gap=0.65)

slide = base_slide("06  |  Terraform, Kubernetes, and GitOps", "Infrastructure creates the cluster; GitOps continuously reconciles workloads")
card(slide, 0.8, 1.65, 3.6, 1.55, "Terraform → AWS", "VPC · public/private subnets · NAT · EKS · managed node group", ORANGE)
card(slide, 4.85, 1.65, 3.6, 1.55, "Argo CD → Helm", "Watch the Git path and sync the declared Helm release.", GREEN)
card(slide, 8.9, 1.65, 3.6, 1.55, "Kubernetes resources", "UI/API Deployments · Services · Ingress · PostgreSQL PVC · HPA", BLUE)
connector(slide, 4.4, 2.43, 4.82, 2.43)
connector(slide, 8.45, 2.43, 8.87, 2.43)
bullet_list(slide, [
    "Readiness protects traffic routing; liveness detects a stuck process.",
    "HPA needs CPU resource requests and a working metrics provider.",
    "A local cluster is useful for learning, but does not prove AWS EKS was deployed."
], y=4.0, size=16, gap=0.66)

slide = base_slide("07  |  Observability", "Measure user-visible behavior and service health")
card(slide, 0.8, 1.7, 3.6, 1.7, "FastAPI /metrics", "HTTP request totals, status codes, and duration histograms.", BLUE)
card(slide, 4.85, 1.7, 3.6, 1.7, "Prometheus", "Scrapes the app through local config or Kubernetes ServiceMonitor.", ORANGE)
card(slide, 8.9, 1.7, 3.6, 1.7, "Grafana", "Visualizes throughput, p95 latency, errors, and activity.", GREEN)
connector(slide, 4.4, 2.55, 4.82, 2.55)
connector(slide, 8.45, 2.55, 8.87, 2.55)
bullet_list(slide, [
    "Check the Prometheus Targets page: the service should be UP.",
    "Generate controlled API traffic and watch time-series panels update.",
    "Metrics explain traffic and behavior; health endpoints answer different probe questions."
], y=4.25, size=16, gap=0.66)

slide = base_slide("08  |  Evidence, demo, and lessons", "Screenshots and recording must show real command output and deployed services")
bullet_list(slide, [
    "Capture terminal commands with your actual prompt/username and host visible.",
    "Show CI, Trivy, GHCR, Terraform/AWS, Argo CD, Kubernetes/Ingress, Prometheus, and Grafana.",
    "Keep a screenshot index and match each image to a genuine command or UI page.",
    "Never expose credentials; destroy billable infrastructure when the demonstration ends.",
    "Explain what ran locally versus in AWS, and call out remaining limitations honestly."
], x=0.95, y=1.75, w=11.4, size=17, gap=0.79)
textbox(slide, 1.0, 6.0, 11.2, 0.42, "Key lesson: a reliable delivery path connects source, tests, artifacts, infrastructure, deployment, and observability.", 14, BLUE, True)

presentation.save(OUTPUT)
print(f"Created {OUTPUT}")
